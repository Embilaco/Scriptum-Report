#!/usr/bin/env python3
# coding: utf-8
#
# part of:
#   S C R I P T U M 
#

###################
# MODULE _docx.reportDocx
# PROVIDES 
#   class ManagedDocx - manage a full word docx-file and combines all other modules in the package
#                       main entry for word-documents

#

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from .paragraphs.element import renderedText
from .section import Sections
from ..rdf.tasks.report_task import GLOBAL_ROOT, ReportTask
from ..tag import Tag, puretagOf

import os

from ..version import version

class ManagedDocx:
    def __init__(self, document: str, debug=False):
        self.document_name = document
        # open and store the Document as given by docx:
        self.document = Document(document)
        self._debug=debug
        
        ## evaluate structure of document and store it
        #self.elements = []

        # in that simple command all work is done to read and organize the document!
        self.sections = Sections(self.document.sections)
        # at least there is always one section

        # collect all errors and warnings
        errors = []
        warnings = []
        for s in self.sections:
            errors += s.error
            warnings += s.warning

        # expose collected diagnostics so callers can inspect them without
        # parsing stdout
        self.errors = errors
        self.warnings = warnings

        # blueprints that left the tree when their first clone was placed --
        # the pruning pass deletes them at the end (typesetting resets it)
        self.spentBlueprints = []

        # instances refused because the block they repeat is not a blueprint;
        # their fills are skipped in silence rather than warned about again
        self.refusedInstances = set()

        if warnings:
            print('There are warnings in the template, outcome might be not as expected:')
            print('  '+'\n  '.join(warnings))

        if errors:
            print('There are errors in the template, further action is needed:')
            print('  '+'\n  '.join(errors))
            print('Please fix prior to proceeding with that template')
            print('PROCESSING STOPPED')
        else:
            # get usable templates
            self.templates = self.sections.templates

            self.toc = self.findTableOfContents()
         
    def apply(self, what, task:ReportTask) -> None:
        """change the tag given as list == path or as string == tag
        
        within this step we do fill and replace content
        """

        #print('\nstart to apply:', what, task)

        if type(what) == str:
            # this may return many nad is used in global search only!
            found = self.sections.findGlobal(what)
        else:
            # this will return one and is used in each single task
            root = self.sections.byName(what[0])
            if not root:
                print(f'WARNING: cannot find section {what[0]!r}')
                return
            parent = root.addressbook.get('.'.join(what[:-1]),None)
            if not parent and '.'.join(what[:-1]) in self.refusedInstances:
                # the instance was refused with its own message; its fills
                # have nowhere to land, which is not news
                return
            if parent:
                # A flagged block is a blueprint, never content. The fill
                # skips it and keeps scanning -- an add's clone is appended
                # *behind* the blueprint in structure order, so this is what
                # lets the ::1 fill reach the instance the add placed at its
                # marker instead of writing into the doomed blueprint.
                found = parent.findExact(what, warn=False, skipBlueprints=True)
                if not found:
                    # Nothing real holds the address. Either it is a
                    # blueprint's first use -- then the ladder rule applies
                    # to it too: clone it where it stands, fill the clone --
                    # or there is truly nothing, which warns as before.
                    found = parent.findExact(what)
                    if found and self._isBlueprintBlock(found[0]):
                        found = [self._cloneBlueprintForFill(found[0], parent,
                                                             what, root)]
                #print('F',found)
            else:
                print(f'WARNING: cannot find parent structure {(".".join(what[:-1]))}')
                return

        #print('found for apply ...', found)
        value = task.value

        # do something special for tables and images when they are standalone
        if value.subtype in [ 'table', 'image' ]:
            # DocImageBlockElement or DocTableBlockElement
            for t,e in found:
                e.fill(task)

        else:
            # do something general for paragraphs (even in tables) when there are tags
            # in case of non-existing files above, write a message if possible
            # otherwise find inline tables, images, images in inline tables etc. as well
            for t,e in found:
                if hasattr(e, 'subtype') and e.subtype == 'text':
                    # DocTextBlockElement
                    e.fill(task)
                elif type(t) == str:
                    print(f'ERROR: FIX - no fill on {t} {e}')
                    continue
                elif not t or t.burned: 
                    continue
                else:
                    # The tag as the document spells it, not the canonical
                    # address: findExact has already picked out this exact
                    # tag, and the text to match on is what is written
                    # there. The instance rides as an `id` argument now,
                    # so a clone's tag still reads `text:description`.
                    self.fillGeneric(t.puretag,t,e,value,
                                     task.actions if task.modified else None)

    def _nothingToApply(self, address):
        """Why an address found nothing, in one line.

        A container address is **positional**: ``findExact`` walks the parent
        that was named and nothing else, so the commonest way to miss is to
        name a block under the wrong parent. Where the blueprint really stands
        is knowable, and saying it turns a report that something is missing
        into one that says what to fix -- which matters here more than
        anywhere, because the entry's own fills each report their missing
        parent afterwards and none of *those* lines can explain the cause.

        Read straight from ``sections.templates`` rather than through
        ``findTemplate``: that lookup warns about an ambiguity, which is the
        last thing a message explaining a *different* mistake should add to
        the output, and a nested blueprint appears in the list once per level
        it is reached through, so the places are deduplicated here.
        """
        wanted = puretagOf(address[-1])
        named = '.'.join(address[:-1])
        parent = puretagOf(address[-2]) if len(address) > 1 else ''

        elsewhere = []
        for _tag, element in self.sections.templates:
            if not element.path or puretagOf(element.path[-1]) != wanted:
                continue
            stands = '.'.join(element.path[:-1])
            if stands and stands != named and stands not in elsewhere:
                elsewhere.append(stands)

        line = f'WARNING: Nothing to apply at {".".join(address)}'
        if not elsewhere:
            return f'{line} - the template holds no {wanted!r} to apply there'
        return (f'{line} - {parent!r} holds no {wanted!r}; it stands under '
                f'{" and ".join(elsewhere)}, and an address is positional')

    @staticmethod
    def _namesakeInPlace(parent, address):
        """Whether *parent* still holds a block under this address's name.

        Tells "the block is there, just not flagged" apart from "that name
        is nowhere in the template", which are the same miss to
        ``findTemplate`` and different mistakes to the author.

        ``parent.structure`` is the right place to look: a blueprint *leaves*
        it when its first instance is cloned, so a block still standing under
        that name is content -- either an unflagged block filled in place, or
        the clone of a blueprint, and in the second case ``findTemplate``
        would have found the blueprint and we would never be asked.
        """
        wanted = puretagOf(address[-1])
        return any(t in ('struct', 'table', 'image', 'text')
                   and e.path and puretagOf(e.path[-1]) == wanted
                   for t, e in parent.structure)

    @staticmethod
    def _isBlueprintBlock(entry):
        """A block entry whose tag says ``template`` -- a blueprint.

        Only block entries count: a simple tag carrying ``isTemplate``
        (inside ``section:template``) has no structure to clone.
        """
        t, e = entry
        return (t in ('text', 'struct', 'table', 'image')
                and getattr(e, 'isTemplate', False))

    def _cloneBlueprintForFill(self, entry, parent, address, root):
        """First use of an in-content blueprint: clone it, fill the clone.

        The ladder rule (typesetting stage 1) extended to ``table:``/
        ``image:``/``text:`` blocks: the clone lands exactly where the
        blueprint stands (before its opening paragraph), the blueprint
        leaves the parent's structure so no later lookup resolves to it,
        and it is remembered as spent for the pruning pass. A blueprint
        that was never through ``getTemplates()`` -- nested inside a
        clone -- gets its deepcopy here.
        """
        t, element = entry
        if not hasattr(element, 'deepcopy'):
            element.createTemplate()
        firstElement = element.structure[0][1]
        parent.structure = [(pt, pe) for pt, pe in parent.structure
                            if pe is not element]
        self.spentBlueprints.append(element)
        clone = element.copy(firstElement, parent=parent,
                             newpath=address[:-1], newname=address[-1],
                             section=root)
        return (t, clone)

    def fillGeneric(self, target: str, tag: Tag, elem, value, actions=None):
        """always do that loop for paragraph type elements and text value types"""
        #obj = determineElement(elem)
        value.load()
        #print('fillGeneric 1a:',elem,target,tag.puretag,value)
        # What the value says beyond str(value) -- shared with the
        # placeholders of a text block, so a file renders the same wherever it
        # is written. It goes *into* the replacement rather than over the run
        # afterwards: `found.text = ...` overwrites the whole run, which ate
        # the words either side of a tag that shared one with them.
        text = renderedText(tag, value)
        found = elem.replaceTagInAll(target, str(value) if text is None else text)
        #print('fillGeneric 1b:',found, elem, value.tostring)

        # a color modifier paints the font of the text this fill wrote --
        # only where a run came back, so template text is never touched
        if actions and hasattr(found, 'font'):
            colour = actions.get('color')
            if colour is not None and colour.type == 'color':
                found.font.color.rgb = RGBColor.from_string(colour.object.for_docx)

    def structure(self, rdf):
        """Expand the ladder and fill nothing: a document to *read*, not to ship.

        Every instance the report document asks for is created and placed, and
        then everything else is left alone, so the saved file shows the shape
        the run is about to fill:

        * each clone carries its instance number as the document addresses it
          -- ``<subsection:content id=1>``, ``<subsection:content id=2>`` --
          which is the thing hardest to picture from the document alone;
        * a blueprint that was **not** used still stands, still saying
          ``template``, so an unused one is visible as unused and *where* it
          was unused -- inside which clone of its parent;
        * every marker add sits in front of the ``<marker:content/>`` it
          belongs to, already in its place;
        * no value is written, so every fill tag is still readable as the
          address it is.

        It is the seven switches of :meth:`typesetting` in the one combination
        that answers "what did it think I meant?" -- all of them off but the
        add-and-copy stage. Nothing is cleaned, nothing is pruned, the
        template section stays, and the document properties are not stamped:
        the file is a diagnostic and is not a report.

        Save it as usual::

            managed = Scriptum.ManagedDocx('template.docx')
            managed.structure(rdf)
            managed.save('structure.docx')

        **Word only.** ``ManagedPptx`` has no counterpart, and not by
        oversight: a slide is created inside the fill pass (``add_slide`` from
        ``applyTask``), so a PowerPoint run with the fill switched off
        produces an empty deck rather than a structure to look at.
        """
        self.typesetting(rdf,
                         addcopy=True,          # the only stage that runs
                         directfill=False,      # leave the fill tags readable
                         globalfill=False,      # and the header/footer ones
                         cleanup=False,         # keep the tags as written
                         removetemplate=False,  # show the blueprints too
                         cleardust=False,       # keep the opening paragraphs
                         setproperties=False)   # this is not a report

    def typesetting(self, rdf, 
                    addcopy=True, 
                    directfill=True,
                    globalfill=True,
                    cleanup=True,
                    removetemplate=True,
                    cleardust=True,
                    setproperties=True
                    ):
        """the final marriage between document and rdf and content
        
        usually all options are True, setting them to False will skip the step during typesetting
        however, this may cause some unintendent results, so use it with care and for debugging only

        * addcopy - apply all add and copy operations
        * directfill - fill all content that is directly addressed
        * globalfill - fill all global addressed content
        * cleanup - run cleanup - remove all remaining xml-syntax from document
        * removetemplate - remove the template section
        * cleardust - remove paragraphs initially marked for deletion
        * setproperties - set document properties
        """
        print('check consistency')

        # Blueprints that were cloned in place of being filled. They leave the
        # tree when their first instance is cloned, so the pruning pass at the
        # end cannot find them by walking it; they are remembered here instead.
        self.spentBlueprints = []
        self.refusedInstances = set()

        if not addcopy:
            print('   SKIP: add and copy new paragraphs and more ...')
        else:
            print('   add and copy new paragraphs and more ...')

            #lasttask = None
            for t in rdf.tasks:
                if t.path[0] == GLOBAL_ROOT: continue # apply the global tasks at the end
                if t.modified: # the modification tells me if I have to add or copy templates
                    #print('\nto %s  *******************\n'%t.what, 
                    #      f'{t.path} - {t.myAddress} - {t.where}')

                    root = self.sections.byName(t.myAddress[0])

                    if not root:
                        print(f'ERROR: cannot find section: {t.myAddress[0]}')
                        continue

                    if t.what == 'apply':
                        if len(t.myAddress) == 1:
                            # A top-level section. There is no parent to
                            # claim a subAnchor from, and the section is
                            # already in the document, so this is a no-op.
                            continue
                        # 'apply' is used on structures and sections that already exist in the document
                        #
                        # will just fill the already existing content later
                        # but we will remove any subAnchor from parent if it exists
                        # this happens always befor we do 'copy' a new section below!
                        parent = root.addressbook.get('.'.join(t.myAddress[:-1]),None)
                        if not parent:
                            print(f'WARNING: cannot find parent structure '
                                  f'{(".".join(t.myAddress[:-1]))}')
                            continue

                        # warn=False: findExact's own miss says the same thing
                        # in its own address notation, and the line below says
                        # it better. This is the only caller that knows what
                        # the miss means.
                        struct = parent.findExact(t.myAddress, warn=False)
                        if not struct:
                            print(self._nothingToApply(t.myAddress))
                            continue

                        element = struct[0][1]
                        firstElement = element.structure[0][1]

                        if element.isTemplate:
                            # A block whose tag says `template` is a blueprint,
                            # and every instance of it is a clone -- the first
                            # included. The clone goes exactly where the
                            # blueprint stands (before its opening paragraph),
                            # which keeps document order whatever order the
                            # data names things in; the blueprint itself is
                            # pruned at the end. Word now agrees with
                            # PowerPoint, which always copies.
                            #
                            # The blueprint leaves the parent's structure
                            # first: findExact takes the first match, and the
                            # fills to come must reach the clone, not a block
                            # that is about to be deleted. A blueprint nested
                            # inside a clone was never through getTemplates(),
                            # so it gets its deepcopy here.
                            if not hasattr(element, 'deepcopy'):
                                element.createTemplate()
                            parent.structure = [(pt, pe) for pt, pe in parent.structure
                                                if pe is not element]
                            self.spentBlueprints.append(element)
                            element.copy(firstElement, parent=parent,
                                         newpath=t.myAddress[:-1],
                                         newname=t.myAddress[-1], section=root)
                            # This clone sits directly before the blueprint's
                            # opening paragraph, so that paragraph is the gap
                            # just behind it -- where instance 2 belongs, and
                            # not wherever the next unused sibling stands.
                            # See StructuredElement.followInstance.
                            parent.followInstance(puretagOf(t.myAddress[-1]),
                                                  firstElement)

                        # Claims this child *and everything ahead of it*, so a
                        # later clone cannot be inserted upstream of the block
                        # it follows -- see StructuredElement.claimSubAnchor.
                        parent.claimSubAnchor(firstElement)

                        continue

                    if t.what == 'add':
                        # 'add' is used when we add by '+' new content from templates after an @anchor
                        parent = root.addressbook.get('.'.join(t.myAddress[:-1]),None)
                        if not parent:
                            print(
                                f'WARNING: cannot find parent structure {(".".join(t.myAddress[:-1]))}'
                            )
                            continue
                        where = parent.findExact(t.myAddress[:-1]+[t.where])
                        if not where:
                            print(
                                f'WARNING: No place to add found: {t.myAddress} {t.where} {where}'
                            )
                            continue
                        
                        anchor = where[0][1]
                        tpl = self.sections.findTemplate(t.target)
                        #print('   add tpl and anchor 0', t.myAddress, where, anchor, tpl)

                        if tpl:
                            #print('   add tpl and anchor 1', parent)
                            newElement = tpl.copy(anchor, parent=parent, newpath=t.myAddress[:-1], newname=t.myAddress[-1], section=root)
                            

                    elif t.what == 'copy':
                        # 'copy' is used in any case we need to duplicate e.g. a section while the existing one is already in place
                        #print('addressbook is:')
                        #for k,v in root.addressbook.items():
                        #    print(' k,v',k,v)
                        #print('looking for',t.myAddress)
                        parent = root.addressbook.get('.'.join(t.myAddress[:-1]), None)
                        #print('parent is', parent)
                        if not parent or not parent.anchor:
                            print(f'WARNING: No place to copy found: {(".".join(t.myAddress[:-1]))}')
                            continue
                        
                        tpl = self.sections.findTemplate(t.path, warn=False)

                        if not tpl:
                            # Repeating a block requires the `template`
                            # argument: without it there is no blueprint to
                            # clone from. Say which block and what to add --
                            # the bare 'no such template' sent the author
                            # looking for a missing name rather than a
                            # missing argument, and the fills inside the
                            # instance that was never made then warned again,
                            # once each, about a parent structure nobody had
                            # asked to be missing.
                            name = puretagOf(t.myAddress[-1])
                            if self._namesakeInPlace(parent, t.myAddress):
                                print(f'WARNING: cannot repeat {name!r}: the '
                                      f'block is in the template but its tag '
                                      f"does not carry the 'template' "
                                      f'argument, so there is nothing to '
                                      f'clone - write <{name} template> to '
                                      f'repeat it')
                            else:
                                print(f'WARNING: No such template in document: {t.path}')
                            self.refusedInstances.add('.'.join(t.myAddress))
                            continue

                        # Right behind the instance this one repeats, if that
                        # instance has been placed -- the template's own prose
                        # between two blocks belongs after all of them, and
                        # the next unclaimed sibling is on the far side of it.
                        anchor = parent.followOnAnchor(puretagOf(t.myAddress[-1]))
                        if anchor is None:
                            if parent.subAnchors:
                                #print('subAnchors')
                                #for a in parent.subAnchors:
                                #    print('   ',a.thing.text)
                                anchor = parent.subAnchors[0]
                            else:
                                #print('anc',parent.anchor)
                                anchor = parent.anchor

                        tpl.copy(anchor, parent=parent, newpath=t.myAddress[:-1],
                                 newname=t.myAddress[-1], section=root)

        if not directfill:                
            print('   SKIP: fill the content...')
        else:
            print('   fill the content...')
            for t in rdf.tasks:
                if t.path[0] == GLOBAL_ROOT: continue # apply the global tasks at the end
                #print('\n',t.isCopy, t.path,t.myAddress,t.target,t.value)
                if t.target:
                    # apply it on path, means: apply it on exact this item
                    #print('  apply to', t.myAddress, t.target, 'v',t.value.object, 'v')
                    self.apply(t.myAddress,t)

        if not globalfill:                
            print('   SKIP: fill the global content...')
        else:
            print('   fill the global content...')
            # apply the global tasks
            for t in rdf.tasks:
                if t.path[0] == GLOBAL_ROOT and t.target:
                    self.apply(t.target,t)

        if not cleanup:                
            print('   SKIP: clean up...')
        else:
            print('   clean up...')
            # finally remove all tags which are not yet "burned"
            for sec in self.sections:
                for t,e in sec:
                    #print('clean1', sec, t)
                    if t in ['text','image','table','struct']:
                        e.clean()
                        continue
                    if not t or t.burned: continue
                    e.replaceTagInAll(t.puretag,'')
                    t.burn()

            # for t,e in self.sections.copies:
            #     if not t or e.anchor or t.burned: continue
            #     #obj = determineElement(e)
            #     e.replaceTagInAll(t.puretag,'')
            #     t.burn()

        if not removetemplate:
            print('   SKIP: remove template section and blueprints...')
        else:
            print('   remove template section and blueprints...')
            self.sections.delete('template')
            self.pruneBlueprints(self.spentBlueprints)

        if not cleardust:                
            print('   SKIP: clearing all the dust...')
        else:
            print('   clearing all the dust...')

            for sec in self.sections:
                for e in sec.markedForDeletion:
                    #print('md',e)
                    e.deleteIfEmpty()

        #self.cleanTableOfContent()

        if not setproperties:                
            print('   SKIP: set properties...')
        else:
            print('   set properties...')

            documenttitle = rdf.settings.documenttitle
            self.document.core_properties.author = f'Scriptum {version}'
            self.document.core_properties.title = documenttitle

        print('done')

    def pruneBlueprints(self, spent):
        """Remove every blueprint from the document, used or not.

        A blueprint is any block whose tag says ``template`` -- a
        section-ladder block or an in-content ``table:``/``image:``/
        ``text:`` block (an open text block outside ``section:template``
        explores as a ``struct``). It is never content: its instances are
        clones, and once they are placed the blueprint has nothing left to
        do -- and an unused one would stay in the finished document
        otherwise, tags cleaned and sample text intact, which is what used
        to happen (for in-content blocks until the ladder rule was extended
        to them).

        Two kinds have to be found two ways. A blueprint that was cloned left
        its parent's structure when that happened, so it is taken from
        ``spent``. One that was never used is still in the tree -- whether it
        came with the template or rides inside a clone, where a nested
        blueprint is copied along with everything else -- so the tree is
        walked for it. Clones themselves do not say ``template`` (numbering
        drops it), so the walk cannot mistake an instance for its blueprint.
        Deleting is a no-op on a detached element, so the overlap between
        the two is harmless, and so is meeting a nested blueprint after its
        parent already went.
        """
        found = list(spent)
        for sec in self.sections:
            if sec.name == 'template':
                continue  # went wholesale, just above
            found += [e for e in sec.iterOnStructures() if e.isTemplate]
        for e in found:
            e.delete(verbose=False)

    def findTableOfContents(self):
        """find the table of contents, other tables untouched for now"""
        b=self.document._body._body
        if (m := b.find(qn('w:sdt'))) != None:
            #print(m)
            try: 
                n=m.find(qn('w:sdtPr')).find(qn('w:docPartObj')).find(qn('w:docPartGallery')).attrib[qn('w:val')]
                #print('n',n)
            except:
                return None
            if n != 'Table of Contents': return None # check for localization?
            try:
                toc = m.find(qn('w:sdtContent'))
                return toc
            except:
                return None
        return None

    def cleanTableOfContent(self):
        """reduce table of content since it will be updated only manually either way
        
        yet not working, TOC broken afterwards
        """
        
        for i,c in enumerate(self.toc.iterchildren()):
            if i < 2: continue
            c.clear()

    def save(self, filename, finish=False, createpdf=False):
        """do a final cleanup and save the result"""
        
        # save from with python-docx
        if filename == self.document_name:
            print('Sorry, overwriting by same name is yet not allowed!')
            return
        
        self.document.save(filename)
        
        if finish:
            if os.name != 'nt':
                print(f'Running on {os.name} will prevent any finishing work...')
            else:
                # win32com is imported at finish time, not at module import:
                # a Windows install without pywin32 keeps a working back end,
                # and only the one step that needs Office reports the gap.
                try:
                    import win32com.client
                except ImportError:
                    print('finishing needs pywin32 -- the Scriptum-Report[windows] '
                          'extra; the document is saved, but unfinished')
                    return

                in_file = os.path.abspath(filename)
                out_file = os.path.abspath(os.path.splitext(os.path.basename(filename))[0]+'.pdf')
                wdFormatPDF = 17

                try:
                    try:
                        word = win32com.client.GetActiveObject('Word.Application')
                        doquit = False
                    except:
                        word = win32com.client.Dispatch('Word.Application')
                        doquit = True
                    try:
                        word.Visible = False
                    except:
                        pass
                    doc = word.Documents.Open(in_file)
                    doc.Fields.Update()
                    try:
                        for i in range(1,doc.TablesOfContents.Count+1):
                            #print(i)
                            doc.TablesOfContents(i).Update()
                    except:
                        pass
                    try:
                        for i in range(1,doc.TablesOfFigures.Count+1):
                            #print(i)
                            doc.TablesOfFigures(i).Update()
                    except:
                        pass
                    doc.Save()
                    if createpdf:
                        doc.SaveAs(out_file, FileFormat=wdFormatPDF)
                    doc.Close(SaveChanges=True)
                    if doquit:
                        word.Quit()
                except Exception as e:
                    print(f'failed to update tables and/or save as PDF\nReason:\n{e}')

